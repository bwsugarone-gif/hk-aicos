# -*- coding: utf-8 -*-
"""
utils/rag_manager.py
HK-AICOS Phase 3.1F — RAG Lite document processing layer.

Local JSON index only. No OCR, no Qdrant, no vector database.
"""

import json
import math
import re
import sys
from datetime import datetime, timezone
from pathlib import Path


_BASE_DIR = Path(__file__).parent.parent
RAG_DIR = _BASE_DIR / "rag_documents"
INDEX_FILE = _BASE_DIR / "rag_index.json"
LEGACY_INDEX_FILE = RAG_DIR / "rag_index.json"

CATEGORIES = [
    "regulations",
    "codes_of_practice",
    "practice_notes",
    "technical_circulars",
    "guidelines",
    "forms_checklists",
    "company_sop",
    "project_docs",
]

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".md", ".txt"}
PRIMARY_SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx"}
CHUNK_MIN = 800
CHUNK_TARGET = 1000
CHUNK_MAX = 1200

CATEGORY_LABELS = {
    "regulations": "法規 / Regulations",
    "codes_of_practice": "實務守則 / Codes of Practice",
    "practice_notes": "實務備考 / Practice Notes",
    "technical_circulars": "技術通告 / Technical Circulars",
    "guidelines": "指引 / Guidelines",
    "forms_checklists": "表格及核查清單 / Forms & Checklists",
    "company_sop": "公司 SOP / Company SOP",
    "project_docs": "項目文件 / Project Documents",
}

_DEPT_KEYWORDS = [
    (["labour", "勞工", "工人安全", "工傷", "工資", "factory", "industrial"], "勞工處"),
    (["bd", "屋宇署", "building department", "buildings ordinance"], "屋宇署"),
    (["fsd", "消防", "fire services", "fire safety"], "消防處"),
    (["emsd", "機電", "electrical", "mechanical", "lift", "escalator"], "EMSD"),
    (["wsd", "水務", "water supply", "waterworks"], "水務署"),
    (["dsd", "渠務", "drainage", "sewerage"], "渠務署"),
    (["hyd", "路政", "highway", "road works"], "路政署"),
    (["td", "運輸", "transport", "traffic"], "運輸署"),
    (["landsd", "地政", "land registry", "lease"], "地政總署"),
    (["epd", "環保", "environmental", "noise", "waste"], "環保署"),
    (["cedd", "geo", "slope", "斜坡", "土力", "geotechnical"], "CEDD / GEO"),
]

_AGENT_QUERY_TERMS = {
    "pm": ["project management", "項目管理", "進度", "協調", "風險"],
    "safety": ["safety", "安全", "高空", "危害", "風險評估"],
    "engineering": ["engineering", "結構", "技術", "施工方法", "工程"],
    "qs": ["cost", "合約", "工程量", "索償", "成本"],
    "accounting": ["payment", "付款", "財務", "發票", "成本"],
    "foreman": ["site", "現場", "施工", "工序", "人手"],
    "material": ["material", "物料", "供應", "採購", "delivery"],
    "surveying": ["survey", "測量", "尺寸", "座標", "放線"],
    "drafting": ["drawing", "圖則", "詳圖", "則師", "revision"],
    "hk_legal": ["legal", "法規", "合規", "條例", "permit"],
    "legal": ["legal", "法規", "合規", "條例", "permit"],
    "translator": ["translation", "翻譯", "文件", "準確", "bilingual"],
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _rel(path: Path) -> str:
    try:
        return str(path.relative_to(_BASE_DIR))
    except Exception:
        return str(path)


def _guess_department(text: str) -> str:
    lower = str(text or "").lower()
    for keywords, dept in _DEPT_KEYWORDS:
        if any(keyword.lower() in lower for keyword in keywords):
            return dept
    return ""


def _guess_language(text: str) -> str:
    value = str(text or "")
    cjk_count = len(re.findall(r"[\u4e00-\u9fff]", value))
    eng_count = len(re.findall(r"[A-Za-z]", value))
    if cjk_count and eng_count:
        return "mixed"
    if cjk_count:
        return "zh-Hant"
    return "en"


def _clean_text(text: str) -> str:
    text = str(text or "").replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _tokenise(text: str) -> list[str]:
    value = str(text or "").lower()
    cjk = re.findall(r"[\u4e00-\u9fff]{2,6}", value)
    eng = re.findall(r"[a-z0-9][a-z0-9_-]{2,}", value)
    return cjk + eng


def _summary(text: str) -> str:
    for line in _clean_text(text).splitlines():
        line = line.strip("#-•* \t")
        if line:
            return line[:160]
    return ""


def _extract_keywords(text: str, max_keywords: int = 30) -> list[str]:
    seen = []
    for token in _tokenise(text):
        if token not in seen:
            seen.append(token)
        if len(seen) >= max_keywords:
            break
    return seen


def _read_text_file(path: Path) -> list[dict]:
    for enc in ("utf-8", "utf-8-sig", "big5", "cp950", "latin-1"):
        try:
            return [{"text": path.read_text(encoding=enc), "section": "text"}]
        except Exception:
            continue
    raise UnicodeDecodeError("unknown", b"", 0, 1, "unsupported text encoding")


def _read_pdf(path: Path) -> tuple[list[dict], str]:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise RuntimeError(f"pypdf unavailable: {exc}") from exc

    reader = PdfReader(str(path))
    sections = []
    for idx, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append({"text": text, "page": idx})

    if not sections:
        return [], "OCR_REQUIRED"

    total_text = "\n".join(section["text"] for section in sections)
    if len(total_text.strip()) < 80 and len(reader.pages) > 0:
        return sections, "OCR_REQUIRED"
    return sections, ""


def _read_docx(path: Path) -> list[dict]:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f"python-docx unavailable: {exc}") from exc

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return [{"text": "\n".join(paragraphs), "section": "document"}]


def _read_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError(f"openpyxl unavailable: {exc}") from exc

    wb = load_workbook(str(path), read_only=True, data_only=True)
    sections = []
    for ws in wb.worksheets:
        lines = []
        for row in ws.iter_rows(values_only=True):
            values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if values:
                lines.append(" | ".join(values))
        if lines:
            sections.append({"text": "\n".join(lines), "sheet": ws.title})
    wb.close()
    return sections


def _extract_sections(path: Path) -> tuple[list[dict], str]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return _read_text_file(path), ""
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path), ""
    if suffix == ".xlsx":
        return _read_xlsx(path), ""
    return [], "UNSUPPORTED_FILE_TYPE"


def _chunk_section(text: str, file_name: str, category: str, chunk_start: int, page="", sheet="", section="") -> list[dict]:
    text = _clean_text(text)
    if not text:
        return []

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    chunks = []
    current = ""
    chunk_id = chunk_start

    def flush():
        nonlocal current, chunk_id
        if not current.strip():
            return
        chunks.append({
            "file_name": file_name,
            "category": category,
            "chunk_id": f"{Path(file_name).stem}-{chunk_id:04d}",
            "page": page,
            "sheet": sheet,
            "section": section,
            "text": current.strip(),
        })
        chunk_id += 1
        current = ""

    for para in paragraphs:
        if len(para) > CHUNK_MAX:
            flush()
            for start in range(0, len(para), CHUNK_TARGET):
                piece = para[start:start + CHUNK_TARGET]
                chunks.append({
                    "file_name": file_name,
                    "category": category,
                    "chunk_id": f"{Path(file_name).stem}-{chunk_id:04d}",
                    "page": page,
                    "sheet": sheet,
                    "section": section,
                    "text": piece.strip(),
                })
                chunk_id += 1
            continue

        candidate = (current + "\n" + para).strip() if current else para
        if len(candidate) > CHUNK_MAX and len(current) >= CHUNK_MIN:
            flush()
            current = para
        else:
            current = candidate

    flush()
    return chunks


def _metadata(path: Path, category: str, text: str, chunk_count: int, status: str = "OK", error: str = "") -> dict:
    combined = f"{path.name}\n{text[:1000]}"
    return {
        "file_name": path.name,
        "file_type": path.suffix.lstrip(".").lower(),
        "category": category,
        "source_department": _guess_department(combined),
        "version_date": "",
        "language": _guess_language(combined),
        "created_at": _now(),
        "summary": _summary(text),
        "chunk_count": chunk_count,
        "path": _rel(path),
        "status": status,
        "error": error,
        "keywords": _extract_keywords(combined),
    }


def load_index() -> dict:
    for path in (INDEX_FILE, LEGACY_INDEX_FILE):
        if path.exists():
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                if "chunks" not in data:
                    data["chunks"] = []
                return data
            except Exception as exc:
                print(f"[rag_manager] WARNING: could not load index {path}: {exc}", file=sys.stderr)
    return {"version": "3.1F", "updated_at": "", "documents": [], "chunks": [], "errors": []}


def save_index(index: dict) -> None:
    index["updated_at"] = _now()
    INDEX_FILE.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")


def scan_documents() -> list[Path]:
    """Scan streamlit-app/rag_documents/ category folders for supported documents."""
    paths = []
    for category in CATEGORIES:
        folder = RAG_DIR / category
        if not folder.exists():
            continue
        for path in sorted(folder.iterdir()):
            if path.is_file() and not path.name.startswith(".") and path.suffix.lower() in SUPPORTED_EXTENSIONS:
                paths.append(path)
    return paths


def build_rag_index() -> dict:
    """Extract text, chunk documents, and write streamlit-app/rag_index.json."""
    documents = []
    chunks = []
    errors = []

    for path in scan_documents():
        category = path.parent.name
        try:
            sections, status = _extract_sections(path)
            doc_text_parts = []
            doc_chunks = []
            chunk_start = 1
            for section_data in sections:
                text = _clean_text(section_data.get("text", ""))
                doc_text_parts.append(text)
                new_chunks = _chunk_section(
                    text=text,
                    file_name=path.name,
                    category=category,
                    chunk_start=chunk_start,
                    page=section_data.get("page", ""),
                    sheet=section_data.get("sheet", ""),
                    section=section_data.get("section", ""),
                )
                chunk_start += len(new_chunks)
                doc_chunks.extend(new_chunks)

            full_text = "\n".join(doc_text_parts)
            if status == "OCR_REQUIRED":
                message = "此文件可能為掃描 PDF，OCR 將於 Phase 3.2 支援。"
                documents.append(_metadata(path, category, full_text, len(doc_chunks), status=status, error=message))
                errors.append({"file_name": path.name, "category": category, "status": status, "message": message})
            else:
                documents.append(_metadata(path, category, full_text, len(doc_chunks), status="OK"))
            chunks.extend(doc_chunks)
        except Exception as exc:
            message = f"{type(exc).__name__}: {exc}"
            documents.append(_metadata(path, category, "", 0, status="ERROR", error=message))
            errors.append({"file_name": path.name, "category": category, "status": "ERROR", "message": message})

    index = {
        "version": "3.1F",
        "updated_at": "",
        "chunk_size": {"min": CHUNK_MIN, "target": CHUNK_TARGET, "max": CHUNK_MAX},
        "documents": documents,
        "chunks": chunks,
        "errors": errors,
    }
    save_index(index)
    print(f"[rag_manager] RAG Lite index built: {len(documents)} docs, {len(chunks)} chunks", file=sys.stderr)
    return index


def _ensure_index() -> dict:
    index = load_index()
    if not index.get("chunks") and not INDEX_FILE.exists():
        return build_rag_index()
    return index


def _score_chunk(chunk: dict, query_tokens: list[str], query: str) -> float:
    text = " ".join([
        chunk.get("file_name", ""),
        chunk.get("category", ""),
        chunk.get("text", ""),
        str(chunk.get("page", "")),
        str(chunk.get("sheet", "")),
        str(chunk.get("section", "")),
    ]).lower()
    if not text:
        return 0.0

    token_hits = sum(1 for token in query_tokens if token and token in text)
    unique_hits = len({token for token in query_tokens if token and token in text})
    exact_bonus = 4 if query and query.lower() in text else 0
    density = unique_hits / math.sqrt(max(len(text), 1))
    return token_hits + unique_hits * 1.5 + exact_bonus + density


def search_rag_context(query: str, top_k: int = 5, category_filter: str = "", department_filter: str = "") -> list[dict]:
    """Keyword + simple semantic overlap search over indexed chunks."""
    index = _ensure_index()
    chunks = index.get("chunks", [])
    docs_by_name = {doc.get("file_name"): doc for doc in index.get("documents", [])}
    tokens = _tokenise(query)
    if not tokens:
        return []

    results = []
    for chunk in chunks:
        doc = docs_by_name.get(chunk.get("file_name"), {})
        if category_filter and chunk.get("category") != category_filter:
            continue
        if department_filter and doc.get("source_department") != department_filter:
            continue
        score = _score_chunk(chunk, tokens, query)
        if score <= 0:
            continue
        results.append({
            **chunk,
            "_score": round(score, 3),
            "source_department": doc.get("source_department", ""),
            "summary": doc.get("summary", ""),
            "language": doc.get("language", ""),
            "status": doc.get("status", "OK"),
        })

    results.sort(key=lambda item: item["_score"], reverse=True)
    return results[:top_k]


def get_relevant_context_for_agents(selected_agents: list, question: str, analysis_type: str, top_k: int = 5) -> str:
    """Return top 3-5 relevant RAG Lite chunks for prompt injection."""
    agent_terms = []
    for agent_id in selected_agents or []:
        agent_terms.extend(_AGENT_QUERY_TERMS.get(str(agent_id), []))
    query = " ".join([analysis_type or "", question or "", " ".join(agent_terms)])
    hits = search_rag_context(query, top_k=max(3, min(top_k, 5)))
    if not hits:
        return ""

    parts = []
    for hit in hits:
        location = ""
        if hit.get("page"):
            location = f" page {hit['page']}"
        elif hit.get("sheet"):
            location = f" sheet {hit['sheet']}"
        elif hit.get("section"):
            location = f" section {hit['section']}"
        dept = f" [{hit['source_department']}]" if hit.get("source_department") else ""
        text = hit.get("text", "")[:1200]
        parts.append(
            f"文件：{hit['file_name']}{dept} ({hit['category']}{location})\n"
            f"相關段落：{text}"
        )
    return "\n\n---\n\n".join(parts)


# ── Backward-compatible wrappers ─────────────────────────────────────────────
def scan_and_index() -> dict:
    return build_rag_index()


def search(query: str, top_k: int = 5, category_filter: str = "", department_filter: str = "") -> list:
    hits = search_rag_context(query, top_k=top_k, category_filter=category_filter, department_filter=department_filter)
    docs = []
    seen = set()
    index = load_index()
    docs_by_name = {doc.get("file_name"): doc for doc in index.get("documents", [])}
    for hit in hits:
        file_name = hit.get("file_name")
        if file_name in seen:
            continue
        seen.add(file_name)
        doc = dict(docs_by_name.get(file_name, {}))
        doc["_score"] = hit.get("_score", 0)
        docs.append(doc)
    return docs[:top_k]


def search_by_regulations(regulation_keys: list, top_k: int = 5) -> list:
    return search(" ".join(regulation_keys or []), top_k=top_k)


def get_document_by_name(file_name: str) -> dict | None:
    index = load_index()
    for doc in index.get("documents", []):
        if doc.get("file_name") == file_name:
            return doc
    return None


def get_all_documents(category: str = "") -> list:
    docs = load_index().get("documents", [])
    if category:
        return [doc for doc in docs if doc.get("category") == category]
    return docs


def get_index_stats() -> dict:
    index = load_index()
    by_category = {}
    for doc in index.get("documents", []):
        category = doc.get("category", "unknown")
        by_category[category] = by_category.get(category, 0) + 1
    return {
        "total": len(index.get("documents", [])),
        "chunks": len(index.get("chunks", [])),
        "errors": len(index.get("errors", [])),
        "updated_at": index.get("updated_at", ""),
        "by_category": by_category,
    }


def read_document(doc: dict, max_chars: int = 2000) -> str:
    file_name = doc.get("file_name", "")
    chunks = [chunk for chunk in load_index().get("chunks", []) if chunk.get("file_name") == file_name]
    if chunks:
        return "\n\n".join(chunk.get("text", "") for chunk in chunks)[:max_chars]
    if doc.get("status") == "OCR_REQUIRED":
        return "此文件可能為掃描 PDF，OCR 將於 Phase 3.2 支援。"
    return doc.get("summary", "")


if __name__ == "__main__":
    idx = build_rag_index()
    stats = get_index_stats()
    print(f"Index saved to: {INDEX_FILE}")
    print(f"Documents: {stats['total']} | Chunks: {stats['chunks']} | Errors: {stats['errors']}")
    if len(sys.argv) > 1:
        q = " ".join(sys.argv[1:])
        for hit in search_rag_context(q, top_k=5):
            print(f"[{hit['_score']}] {hit['file_name']} / {hit['chunk_id']}: {hit['text'][:80]}")
