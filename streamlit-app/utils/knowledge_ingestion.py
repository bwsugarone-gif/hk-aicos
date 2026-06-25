"""Knowledge / PDF / document ingestion into knowledge sources + RAG (Phase 6.4).

Turns an uploaded PDF / TXT / MD / DOCX (XLSX: metadata only) into:

* a persisted :class:`KnowledgeSource` (append-only ``data/ingested_knowledge.jsonl``)
* deterministic RAG chunks with page references (see :mod:`utils.rag_persistence`)
* a file-registry entry linking the original file metadata

Scanned / image-only PDFs degrade to a *metadata-only* knowledge source with a
clear Traditional-Chinese warning — no heavy OCR dependency is added.
"""

from __future__ import annotations

import hashlib
import json
import os
import threading
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .knowledge_models import KnowledgeSource
from .pdf_text_extractor import (
    METADATA_ONLY_WARNING,
    STATUS_METADATA_ONLY,
    STATUS_SELECTABLE,
    ExtractionResult,
    extract_pdf_pages,
    extract_plain_text,
    result_from_text,
)
from .rag_persistence import (
    DEFAULT_INGESTED_RAG_PATH,
    append_ingested_chunks,
    build_chunks_from_pages,
)
from .rag_models import RagChunk


_DATA_DIR = Path(__file__).resolve().parents[1] / "data"
DEFAULT_INGESTED_KNOWLEDGE_PATH = _DATA_DIR / "ingested_knowledge.jsonl"
_LOCK = threading.Lock()

_PDF_EXTS = {".pdf"}
_TEXT_EXTS = {".txt", ".md", ".rst"}
_DOCX_EXTS = {".docx"}
_XLSX_EXTS = {".xlsx"}


@dataclass
class IngestionResult:
    source: KnowledgeSource
    chunks: list[RagChunk] = field(default_factory=list)
    status: str = STATUS_METADATA_ONLY
    extracted_chars: int = 0
    pages_extracted: int = 0
    chunk_count: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def source_id(self) -> str:
        return self.source.source_id

    @property
    def metadata_only(self) -> bool:
        return self.status != STATUS_SELECTABLE or not self.chunks

    def to_public_dict(self) -> dict[str, Any]:
        """UI-safe summary: no local paths, no raw JSON, no secrets."""
        return {
            "source_id": self.source.source_id,
            "title": self.source.title,
            "status": self.status,
            "extracted_chars": self.extracted_chars,
            "pages_extracted": self.pages_extracted,
            "chunk_count": self.chunk_count,
            "warnings": list(self.warnings),
            "metadata_only": self.metadata_only,
        }


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _append_source(path: Path, source: KnowledgeSource) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with _LOCK:
        with path.open("a", encoding="utf-8", newline="\n") as handle:
            handle.write(json.dumps(source.to_dict(), ensure_ascii=False, separators=(",", ":")) + "\n")
            handle.flush()
            os.fsync(handle.fileno())


def _new_source_id(seed: str) -> str:
    digest = hashlib.sha256(f"{seed}:{_now()}".encode("utf-8", errors="ignore")).hexdigest()[:20]
    return f"uksrc_{digest}"


def read_ingested_knowledge_sources(
    *,
    path: str | Path = DEFAULT_INGESTED_KNOWLEDGE_PATH,
) -> list[KnowledgeSource]:
    """Latest-wins read of ingested knowledge sources (newest first)."""
    file_path = Path(path)
    if not file_path.exists():
        return []
    latest: dict[str, KnowledgeSource] = {}
    try:
        lines = file_path.read_text(encoding="utf-8", errors="replace").splitlines()
    except OSError:
        return []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        try:
            source = KnowledgeSource.from_dict(json.loads(line))
        except (json.JSONDecodeError, TypeError, ValueError, KeyError):
            continue
        if source.source_id:
            latest[source.source_id] = source
    return sorted(latest.values(), key=lambda item: item.last_indexed_at, reverse=True)


def _extract(path: Path) -> tuple[ExtractionResult, list[str]]:
    """Dispatch extraction by extension; return (result, notes)."""
    ext = path.suffix.lower()
    notes: list[str] = []
    if ext in _PDF_EXTS:
        return extract_pdf_pages(path), notes
    if ext in _TEXT_EXTS:
        return extract_plain_text(path), notes
    if ext in _DOCX_EXTS:
        try:
            from docx import Document

            doc = Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    line = " | ".join(dict.fromkeys(c for c in cells if c))
                    if line:
                        text += "\n" + line
            return result_from_text(text), notes
        except ImportError:
            notes.append("未安裝 Word 讀取元件，已建立 metadata-only 知識來源。")
            return ExtractionResult(status=STATUS_METADATA_ONLY, warning=METADATA_ONLY_WARNING), notes
        except Exception:
            notes.append("無法讀取 Word 內容，已建立 metadata-only 知識來源。")
            return ExtractionResult(status=STATUS_METADATA_ONLY, warning=METADATA_ONLY_WARNING), notes
    if ext in _XLSX_EXTS:
        notes.append("Excel 僅建立 metadata 知識來源（不抽取全文）。")
        return ExtractionResult(status=STATUS_METADATA_ONLY, warning=METADATA_ONLY_WARNING), notes
    notes.append("不支援的檔案類型，已建立 metadata-only 知識來源。")
    return ExtractionResult(status=STATUS_METADATA_ONLY, warning=METADATA_ONLY_WARNING), notes


def ingest_document(
    *,
    file_path: str | Path,
    original_file_name: str | None = None,
    project_ref: str | None = None,
    tags: list[str] | None = None,
    description: str = "",
    trust_level: str = "internal",
    persist: bool = True,
    register_in_file_registry: bool = True,
    knowledge_path: str | Path = DEFAULT_INGESTED_KNOWLEDGE_PATH,
    rag_path: str | Path = DEFAULT_INGESTED_RAG_PATH,
    registry_path: Any = None,
) -> IngestionResult:
    """Ingest one uploaded document into a knowledge source + RAG chunks."""
    path = Path(file_path)
    name = original_file_name or path.name
    result, notes = _extract(path)
    return _finalize_ingestion(
        extraction=result,
        notes=notes,
        original_file_name=name,
        project_ref=project_ref,
        tags=tags,
        description=description,
        trust_level=trust_level,
        local_runtime_path=str(path),
        persist=persist,
        register_in_file_registry=register_in_file_registry,
        knowledge_path=knowledge_path,
        rag_path=rag_path,
        registry_path=registry_path,
    )


def ingest_text(
    *,
    text: str,
    title: str,
    original_file_name: str | None = None,
    project_ref: str | None = None,
    tags: list[str] | None = None,
    description: str = "",
    trust_level: str = "internal",
    persist: bool = True,
    knowledge_path: str | Path = DEFAULT_INGESTED_KNOWLEDGE_PATH,
    rag_path: str | Path = DEFAULT_INGESTED_RAG_PATH,
) -> IngestionResult:
    """Ingest raw text (no file) into a knowledge source + RAG chunks."""
    extraction = result_from_text(text)
    return _finalize_ingestion(
        extraction=extraction,
        notes=[],
        original_file_name=original_file_name or f"{title}.txt",
        title=title,
        project_ref=project_ref,
        tags=tags,
        description=description,
        trust_level=trust_level,
        local_runtime_path=None,
        persist=persist,
        register_in_file_registry=False,
        knowledge_path=knowledge_path,
        rag_path=rag_path,
        registry_path=None,
    )


def _finalize_ingestion(
    *,
    extraction: ExtractionResult,
    notes: list[str],
    original_file_name: str,
    project_ref: str | None,
    tags: list[str] | None,
    description: str,
    trust_level: str,
    local_runtime_path: str | None,
    persist: bool,
    register_in_file_registry: bool,
    knowledge_path: str | Path,
    rag_path: str | Path,
    registry_path: Any,
    title: str | None = None,
) -> IngestionResult:
    now = _now()
    tags = list(dict.fromkeys(tags or []))
    source_title = title or Path(original_file_name).stem.replace("_", " ").replace("-", " ")[:160] or "未命名知識來源"
    source_id = _new_source_id(original_file_name)
    has_text = extraction.status == STATUS_SELECTABLE and extraction.pages

    warnings = list(notes)
    if not has_text and extraction.warning:
        warnings.append(extraction.warning)

    if has_text:
        summary = (description or extraction.full_text[:600]).strip()
        source_type = "uploaded_document"
    else:
        summary = (description or f"{Path(original_file_name).suffix.upper().lstrip('.')} 文件；僅建立 metadata 知識來源。").strip()
        source_type = "uploaded_document_metadata_only"

    source = KnowledgeSource(
        source_id=source_id,
        title=source_title,
        source_type=source_type,
        path_or_url=None,  # never expose local path as a knowledge URL
        trust_level=trust_level,
        jurisdiction="HK",
        trade_tags=[],
        topic_tags=tags,
        summary=summary[:1200],
        extracted_refs=[],
        last_indexed_at=now,
    )

    chunks: list[RagChunk] = []
    if has_text:
        chunks = build_chunks_from_pages(
            source_id=source_id,
            source_title=source_title,
            source_file_name=original_file_name,
            pages=extraction.pages,
            project_ref=project_ref,
            source_type="uploaded_document",
            trust_level=trust_level,
            tags=tags,
        )

    if persist:
        _append_source(Path(knowledge_path), source)
        if chunks:
            append_ingested_chunks(chunks, path=rag_path)
        if register_in_file_registry and local_runtime_path:
            _register_file(
                original_file_name=original_file_name,
                project_ref=project_ref,
                local_runtime_path=local_runtime_path,
                source_id=source_id,
                tags=tags,
                has_text=bool(has_text),
                registry_path=registry_path,
            )

    return IngestionResult(
        source=source,
        chunks=chunks,
        status=extraction.status,
        extracted_chars=extraction.total_chars,
        pages_extracted=len(extraction.pages),
        chunk_count=len(chunks),
        warnings=warnings,
    )


def _register_file(
    *,
    original_file_name: str,
    project_ref: str | None,
    local_runtime_path: str | None,
    source_id: str,
    tags: list[str],
    has_text: bool,
    registry_path: Any,
) -> None:
    try:
        from .file_registry import register_file
        from .file_storage import build_file_metadata, guess_file_type

        file_type = guess_file_type(original_file_name, source_module="knowledge_ingestion")
        payload = build_file_metadata(
            original_file_name=original_file_name,
            source_module="knowledge_ingestion",
            project_ref=project_ref,
            local_runtime_path=local_runtime_path,
            file_type=file_type,
            tags=list(dict.fromkeys(["知識來源", *tags])),
            linked_knowledge_source_id=source_id,
            metadata={"has_extracted_text": bool(has_text)},
        )
        kwargs = {"path": registry_path} if registry_path is not None else {}
        register_file(payload, **kwargs)
    except Exception:
        return
